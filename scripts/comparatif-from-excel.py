"""Régénère src/lib/comparatif-data.ts depuis l'Excel de veille concurrentielle.

Usage : python scripts/comparatif-from-excel.py <Comparatif_fonctionnalites_ImmoTopia_3_concurrents_v2.xlsx>
Feuilles attendues : « Comparatif » (en-tête ligne 7, données à partir de la ligne 8) et « Sources ».
"""
import json
import os
import sys
import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else sys.exit(__doc__)
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src", "lib", "comparatif-data.ts")

STATUS = {"Oui": "oui", "Partiel": "partiel", "ND": "nd", "Abs. doc": "absent", "À vérifier": "verif"}
SLUG = {
    "Socle": "socle", "Pilotage": "pilotage", "Biens & commercial": "biens-commercial",
    "Gestion locative": "gestion-locative", "Portails & service": "portails-service",
    "Finance & documents": "finance-documents", "Syndic": "syndic", "Chantiers & BTP": "chantiers-btp",
    "Patrimoine": "patrimoine", "Communication": "communication", "Intégrations": "integrations",
}

wb = openpyxl.load_workbook(XLSX)
ws = wb["Comparatif"]
rows = []
for r in range(8, ws.max_row + 1):
    v = [ws.cell(r, c).value for c in range(1, 9)]
    if not v[1]:
        continue
    rows.append({
        "domain": SLUG[v[0]], "feature": v[1],
        "s": [STATUS[v[2]], STATUS[v[3]], STATUS[v[4]], STATUS[v[5]]],
        "note": v[6], "sources": v[7].split(),
    })

sources = []
ss = wb["Sources"]
for r in range(7, ss.max_row + 1):
    v = [ss.cell(r, c).value for c in range(1, 7)]
    if v[0]:
        sources.append({"ref": v[0], "editor": v[1], "page": v[2], "url": v[3], "limit": v[5]})

# Le site public ne parle ni de branches ni de fusion.
PUBLIC = [("Branche gestion locative en cours, non fusionnée. ", "En cours de déploiement. "),
          (" (branche en cours, non fusionnée)", " (en cours de déploiement)")]
for row in rows:
    for a, b in PUBLIC:
        row["note"] = row["note"].replace(a, b)
for s in sources:
    if s["ref"].startswith("I"):
        s["page"] = "Documentation fonctionnelle ImmoTopia" + (" — nouveautés en cours de déploiement" if s["ref"] == "I2" else "")
        s["url"] = ""
        s["limit"] = "Fonctions décrites par l'éditeur ImmoTopia."

left = sum(1 for r in rows for s in r["s"] if s == "verif")

def js(o):
    return json.dumps(o, ensure_ascii=False)

lines = [
    "// Fichier généré depuis « Comparatif_fonctionnalites_ImmoTopia_3_concurrents_v2.xlsx » : ne pas modifier à la main.",
    "// Ordre des statuts : ImmoTopia, ChezvousBO, Logestimmo, WIMMO.",
    'import type { CompareRow, CompareSource } from "./comparatif";',
    "",
    "export const compareRows: CompareRow[] = [",
]
for r in rows:
    lines.append(f"  {{ domain: {js(r['domain'])}, feature: {js(r['feature'])}, statuses: {js(r['s'])}, note: {js(r['note'])}, sources: {js(r['sources'])} }},")
lines += ["];", "", "export const compareSources: CompareSource[] = ["]
for s in sources:
    lines.append(f"  {{ ref: {js(s['ref'])}, editor: {js(s['editor'])}, page: {js(s['page'])}, url: {js(s['url'])}, limit: {js(s['limit'])} }},")
lines += ["];", ""]
open(OUT, "w", encoding="utf-8", newline="\n").write("\n".join(lines))
print(len(rows), "lignes,", len(sources), "sources,", left, "statuts encore à vérifier")
